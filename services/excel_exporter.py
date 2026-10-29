from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


class ExcelExporter:
    """Экспорт данных испытания в Excel."""

    @staticmethod
    def export(
        file_path: str,
        exp_data,
        comments: list[tuple]
    ):
        """
        Создаёт Excel-книгу с общими данными и комментариями.

        Args:
            file_path: путь для сохранения XLSX-файла.
            exp_data: объект ExpData с данными испытания.
            comments: список комментариев из DatabaseAccessor.
                      Формат записи:
                      (id, created_at, last_edited, text).
        """

        workbook = Workbook()

        # Лист с общими данными испытания.
        general_sheet = workbook.active
        general_sheet.title = "Общие данные"

        general_sheet.append(["Параметр", "Значение"])

        general_data = [
            ("Номер испытания", exp_data.exp_id),
            ("Изделие", exp_data.product_name),
            ("Исполнитель", exp_data.executor_name),
            ("Дата испытания", exp_data.exp_date),
            ("Описание", exp_data.description),
        ]

        for parameter, value in general_data:
            general_sheet.append([parameter, value])

        # Лист с комментариями.
        comments_sheet = workbook.create_sheet("Комментарии")

        comments_sheet.append([
            "Время создания",
            "Последнее изменение",
            "Комментарий",
        ])

        for comment_id, created_at, last_edited, text in comments:
            comments_sheet.append([
                created_at,
                last_edited,
                text,
            ])

        # Оформление заголовков обоих листов.
        for sheet in (general_sheet, comments_sheet):
            for cell in sheet[1]:
                cell.font = Font(bold=True, color="FFFFFF")
                cell.fill = PatternFill(
                    fill_type="solid",
                    fgColor="4472C4",
                )
                cell.alignment = Alignment(
                    horizontal="center",
                    vertical="center",
                    wrap_text=True,
                )

            sheet.freeze_panes = "A2"

        # Выравнивание и перенос текста.
        for row in general_sheet.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = Alignment(
                    vertical="top",
                    wrap_text=True,
                )

        for row in comments_sheet.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = Alignment(
                    vertical="top",
                    wrap_text=True,
                )

        # Подбор ширины столбцов с ограничением максимальной ширины.
        for sheet in (general_sheet, comments_sheet):
            for column in sheet.columns:
                column_letter = get_column_letter(column[0].column)

                max_length = max(
                    len(str(cell.value or ""))
                    for cell in column
                )

                sheet.column_dimensions[column_letter].width = min(
                    max(max_length + 2, 15),
                    60,
                )

        # Сохраняем книгу.
        workbook.save(file_path)